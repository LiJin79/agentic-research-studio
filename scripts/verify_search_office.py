"""一次有边界的真实验收：官方域名搜索、来源持久化和办公文件下载。

会消耗少量已配置的生成模型与搜索额度，只创建并清理自己的验收会话。
不读取其他会话内容，不输出密钥，结果与下载文件保存在 evaluation 中。
"""

import asyncio
import io
import json
import sys
import uuid
from pathlib import Path
from urllib.parse import urlsplit

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

import httpx  # noqa: E402
from app.core.config import get_settings  # noqa: E402
from docx import Document  # noqa: E402
from pptx import Presentation  # noqa: E402


async def main() -> None:
    """通过用户实际部署的 HTTP 接口检查结果，不将模拟测试写成云端成功。"""
    settings = get_settings()
    headers = (
        {"Authorization": f"Bearer {settings.app_access_token}"}
        if settings.app_access_token
        else {}
    )
    output = Path(__file__).resolve().parents[1] / "evaluation" / "search-office"
    output.mkdir(parents=True, exist_ok=True)
    session_id = None
    report = {"mode": "cloud", "checks": {}}
    async with httpx.AsyncClient(
        base_url="http://127.0.0.1:8080/api/v1", headers=headers, timeout=240
    ) as client:
        try:
            response = await client.get("/health/ready")
            response.raise_for_status()
            report["checks"]["ready"] = response.json()["status"] == "ok"
            response = await client.post(
                "/sessions", json={"title": "自动验收-搜索办公-" + uuid.uuid4().hex[:8]}
            )
            response.raise_for_status()
            session_id = response.json()["id"]
            response = await client.post(
                "/chat",
                json={
                    "session_id": session_id,
                    "use_web_search": True,
                    "message": (
                        "请进行网页搜索，搜索词为 site:fastapi.tiangolo.com FastAPI features。"
                        "只使用该域名下官方资料，用三句话介绍 FastAPI 并引用来源。"
                    ),
                },
            )
            response.raise_for_status()
            result = response.json()
            report["route"] = result["metadata"]["route"]
            report["source_urls"] = [s.get("url") for s in result["sources"]]
            report["tool_events"] = result["metadata"].get("tool_events", [])
            report["answer"] = result["answer"]
            report["verification_status"] = result["metadata"].get("verification_status")
            report["checks"]["web_route"] = report["route"] == "web"
            report["checks"]["search_executed"] = any(
                t.get("name") == "search_web" and t.get("success") for t in report["tool_events"]
            )
            report["checks"]["official_sources"] = bool(report["source_urls"]) and all(
                urlsplit(url or "").hostname == "fastapi.tiangolo.com"
                for url in report["source_urls"]
            )
            response = await client.get(f"/sessions/{session_id}/messages")
            response.raise_for_status()
            assistant = next(m for m in response.json()["items"] if m["role"] == "assistant")
            report["checks"]["tools_persisted"] = bool(
                assistant["message_metadata"].get("tool_events")
            )
            report["checks"]["message_id"] = assistant["id"] == result["metadata"].get("message_id")
            for format in ["docx", "pptx"]:
                response = await client.post(
                    "/artifacts",
                    json={
                        "message_id": assistant["id"],
                        "format": format,
                        "title": "FastAPI 官方资料简报",
                    },
                )
                response.raise_for_status()
                filename = output / ("fastapi-report." + format)
                filename.write_bytes(response.content)
                if format == "docx":
                    document = Document(io.BytesIO(response.content))
                    report["checks"]["docx"] = any("FastAPI" in p.text for p in document.paragraphs)
                else:
                    presentation = Presentation(io.BytesIO(response.content))
                    report["checks"]["pptx"] = len(presentation.slides) >= 3
                    report["slide_count"] = len(presentation.slides)
        except Exception as exc:
            report["error"] = str(exc)[:400]
            if isinstance(exc, httpx.HTTPStatusError):
                try:
                    report["detail"] = exc.response.json().get("detail")
                except ValueError:
                    pass
            raise
        finally:
            if session_id:
                response = await client.delete(f"/sessions/{session_id}")
                report["checks"]["own_session_cleaned"] = response.status_code == 204
            (output / "report.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            print(
                json.dumps(
                    {
                        "checks": report["checks"],
                        "error": report.get("error"),
                        "detail": report.get("detail"),
                    },
                    ensure_ascii=True,
                )
            )
    if not all(report["checks"].values()):
        raise RuntimeError("验收条件没有全部通过，请查看 evaluation/search-office/report.json")


if __name__ == "__main__":
    asyncio.run(main())
