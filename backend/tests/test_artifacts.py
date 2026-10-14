"""检查办公技能实际产生可打开的文件，而不是伪造文件路径或成功提示。"""

import asyncio
import io
import json

import pytest
from app.core.exceptions import UpstreamServiceError
from app.llm.factory import MockLLMClient
from app.schemas.artifacts import ArtifactRequest
from app.services.artifact_service import ArtifactService, safe_title
from docx import Document
from pptx import Presentation


class OutlineLLM(MockLLMClient):
    """只返回受控演示大纲，不调用外部 API。"""

    async def generate(self, messages, *, json_mode=False):
        """确定性内容便于检查备注、引用和文本框。"""
        return json.dumps(
            {
                "slides": [
                    {
                        "title": "项目测试信息",
                        "bullets": ["负责人林晓舟 [K1]"],
                        "notes": "讲解项目资料的来源与核验范围。",
                    }
                ]
            },
            ensure_ascii=False,
        )


def source():
    """构造可追溯内部证据。"""
    return {"citation_id": "K1", "title": "knowledge-test.txt", "url": None}


def test_word_has_real_headings_table_and_sources():
    """Word 保留表格、中文与来源信息。"""
    result = asyncio.run(
        ArtifactService(OutlineLLM()).create(
            title="项目报告",
            content="## 测试信息\n负责人是**林晓舟** [K1]。\n\n|项目|值|\n|---|---|\n|编号|4827|",
            format="docx",
            sources=[source()],
            verified=True,
        )
    )
    document = Document(io.BytesIO(result.content))
    assert result.filename == "项目报告.docx"
    assert document.paragraphs[0].style.name == "Title"
    assert (
        document.styles["Title"].element.find(
            "w:pPr/w:pBdr", document.styles["Title"].element.nsmap
        )
        is None
    )
    assert any(p.style.name == "Heading 2" for p in document.paragraphs)
    assert document.tables[0].cell(1, 1).text == "4827"
    assert "knowledge-test.txt" in "\n".join(p.text for p in document.paragraphs)


def test_ppt_has_editable_slides_notes_and_references():
    """PPT 包含封面、内容、来源以及可编辑备注。"""
    result = asyncio.run(
        ArtifactService(OutlineLLM()).create(
            title="项目报告",
            content="负责人林晓舟 [K1]",
            format="pptx",
            sources=[source()],
            verified=True,
        )
    )
    presentation = Presentation(io.BytesIO(result.content))
    assert len(presentation.slides) == 3
    assert "核验范围" in presentation.slides[1].notes_slide.notes_text_frame.text
    assert any(
        "林晓舟" in shape.text for shape in presentation.slides[1].shapes if shape.has_text_frame
    )
    assert any(
        "knowledge-test.txt" in shape.text
        for shape in presentation.slides[2].shapes
        if shape.has_text_frame
    )


def test_invalid_outline_does_not_download_fake_ppt():
    """无效大纲明确失败，不悄悄输出缺内容的文档。"""

    class Invalid(OutlineLLM):
        """模拟大纲 JSON 结构错误。"""

        async def generate(self, messages, *, json_mode=False):
            """返回无效结构。"""
            return '{"slides": []}'

    with pytest.raises(UpstreamServiceError, match="大纲格式无效"):
        asyncio.run(
            ArtifactService(Invalid()).create(
                title="test", content="test", format="pptx", sources=[], verified=False
            )
        )


def test_artifact_boundaries_and_download_filename():
    """禁止任意输出格式及路径注入，过长内容不得被静默截断。"""
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        ArtifactRequest(message_id="00000000-0000-0000-0000-000000000001", format="exe")
    assert "/" not in safe_title("../../report\r\n.pdf")
    with pytest.raises(ValueError, match="超过"):
        asyncio.run(
            ArtifactService(OutlineLLM()).create(
                title="test", content="x" * 18001, format="pptx", sources=[], verified=False
            )
        )


def test_ppt_rejects_model_invented_citations():
    """不能把模型新增的虚假引用编号排入看似可信的办公文件。"""

    class FakeCitation(OutlineLLM):
        """结构合法但引用不存在的大纲。"""

        async def generate(self, messages, *, json_mode=False):
            """提供一个不在实际来源中的编号。"""
            return '{"slides":[{"title":"报告","bullets":["错误来源 [W99]"],"notes":""}]}'

    with pytest.raises(UpstreamServiceError, match="大纲格式无效"):
        asyncio.run(
            ArtifactService(FakeCitation()).create(
                title="test",
                content="负责人 [K1]",
                format="pptx",
                sources=[source()],
                verified=True,
            )
        )
