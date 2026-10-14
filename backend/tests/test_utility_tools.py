"""新增工具的安全、计算与双语模板回归，不消耗云端模型额度。"""

import asyncio
import io
import json
import uuid
from string import Formatter

import httpx
import pytest
from app.agents import prompts, prompts_en
from app.agents.evidence import merge_evidence
from app.agents.tools import network, registry, utility
from app.core.config import get_settings
from app.schemas.chat import SourceReference
from app.vector.parsing import extract_text
from pptx import Presentation


@pytest.mark.parametrize(
    "url",
    [
        "file:///etc/passwd",
        "http://user:pass@example.com",
        "http://example.com:8001",
        "ftp://example.com/file",
    ],
)
def test_reject_unsafe_url(url):
    """不允许危险协议、认证信息与自定义端口。"""
    with pytest.raises(ValueError):
        network.validate_url(url)


def test_reject_private_or_mixed_dns(monkeypatch):
    """解析结果包含任一私网地址即失败，不能挑一个公网地址绕过校验。"""
    monkeypatch.setattr(
        network.socket,
        "getaddrinfo",
        lambda *a, **k: [(2, 1, 6, "", ("93.184.216.34", 443)), (2, 1, 6, "", ("127.0.0.1", 443))],
    )
    with pytest.raises(ValueError):
        asyncio.run(network.public_address("example.com", 443))


def test_pinned_connection_and_redirect_validation(monkeypatch):
    """网络请求使用校验过的 IP，TLS SNI 与 Host 仍使用原始域名。"""
    original_client = httpx.AsyncClient
    seen = []

    async def public(host, port):
        """用固定公网地址隔离真实 DNS。"""
        if host == "localhost":
            raise ValueError("private")
        return "93.184.216.34"

    def handler(request):
        """模拟重定向到内网，观察执行器是否在第二跳前阻止。"""
        seen.append(request)
        return httpx.Response(302, headers={"Location": "http://localhost/private"})

    monkeypatch.setattr(network, "public_address", public)
    monkeypatch.setattr(
        network.httpx,
        "AsyncClient",
        lambda **kw: original_client(**kw, transport=httpx.MockTransport(handler)),
    )
    with pytest.raises(ValueError):
        asyncio.run(network.fetch_bytes("https://example.com/page"))
    assert len(seen) == 1
    assert seen[0].url.host == "93.184.216.34"
    assert seen[0].headers["host"] == "example.com"
    assert seen[0].extensions["sni_hostname"] == "example.com"


def test_csv_calculation_is_real_and_bounded():
    """聚合数字由 DuckDB 计算，不接受模型 SQL 或静默忽略非数字。"""
    body = b"region,sales\nA,10\nA,20\nB,7\n"
    result = utility.calculate_csv(body, "sum", "sales", "region")
    assert result["results"] == [("A", 30.0), ("B", 7.0)]
    assert result["row_count"] == 3
    with pytest.raises(ValueError):
        utility.calculate_csv(body, "COPY input_data TO '/tmp/file'")
    with pytest.raises(ValueError):
        utility.calculate_csv(b"sales\n1\nunknown\n", "avg", "sales")


def test_page_text_replaces_search_summary():
    """同一链接的正文覆盖摘要，仍只有一个稳定引用编号。"""
    base = {"source_type": "web", "url": "https://example.com", "title": "source"}
    results = merge_evidence(
        [
            base | {"snippet": "short", "content_kind": "search_summary"},
            base | {"snippet": "full source text", "content_kind": "page_text"},
        ]
    )
    assert len(results) == 1
    assert results[0]["snippet"] == "full source text"
    assert results[0]["citation_id"] == "W1"


def test_no_unconfigured_or_irrelevant_tools(monkeypatch):
    """缺少 Key 的工具不绑定；指定域名不能绕到其他数据源。"""
    settings = get_settings()
    for name in ["amap_api_key", "qweather_api_key", "context7_api_key"]:
        monkeypatch.setattr(settings, name, "")
    names = {t.name for t in registry.extra_tools("web", "旅游天气 API", [])}
    assert "read_webpage" in names
    assert not names & {"get_weather", "plan_route", "lookup_library_docs"}
    assert {t.name for t in registry.extra_tools("web", "代码论文", ["example.com"])} == {
        "read_webpage"
    }


def test_bilingual_prompt_variables_match():
    """英文档案变量与中文一一对应，后续切换语言不会漏传参数。"""
    for name, value in vars(prompts).items():
        if not name.isupper() or not isinstance(value, str):
            continue
        english = getattr(prompts_en, name)
        fields = {field for _, field, _, _ in Formatter().parse(value) if field}
        english_fields = {field for _, field, _, _ in Formatter().parse(english) if field}
        assert fields == english_fields, name
        value.format(**dict.fromkeys(fields, "test"))
        english.format(**dict.fromkeys(fields, "test"))


def test_read_powerpoint_text():
    """上传 PPT 支持提取实际文字，而不是把图片当已识别文本。"""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Research title"
    slide.placeholders[1].text = "Evidence-based content"
    output = io.BytesIO()
    deck.save(output)
    text = extract_text("source.pptx", output.getvalue())
    assert "Research title" in text
    assert "Evidence-based content" in text


def test_document_tool_result_matches_response_schema(monkeypatch):
    """新文档与计算结果必须能进入聊天响应，不让引用字段在最后保存时才报错。"""
    identifier = str(uuid.uuid4())

    async def fake(document_id):
        """模拟已验证的上传文件，避免真实数据库与文件访问。"""
        return "sales.csv", b"region,sales\nA,10\nA,20\n"

    monkeypatch.setattr(utility, "_document_bytes", fake)
    for result in [
        asyncio.run(utility.read_document(identifier)),
        asyncio.run(utility.analyze_table(identifier, "sum", "sales", "region")),
    ]:
        source = SourceReference.model_validate(merge_evidence(result)[0])
        assert source.document_id == uuid.UUID(identifier)
        assert source.content_kind in {"document_text", "calculation"}
        assert source.chunk_index is None


def test_search_papers_maps_atom(monkeypatch):
    """论文摘要映射保留日期与预印本边界。"""

    async def fake(*args, **kwargs):
        """返回小型 Atom 响应而非真正调用外网。"""
        return (
            b'<feed xmlns="http://www.w3.org/2005/Atom"><entry>'
            b"<id>http://arxiv.org/abs/1234.5678</id><title>Study</title>"
            b"<summary>Abstract</summary><published>2026-01-01</published>"
            b"</entry></feed>",
            "https://export.arxiv.org",
            "application/atom+xml",
        )

    monkeypatch.setattr(utility, "fetch_bytes", fake)
    result = asyncio.run(utility.search_papers("agent"))
    assert result[0]["content_kind"] == "paper_abstract"
    assert result[0]["published_at"] == "2026-01-01"
    assert json.dumps(result)
