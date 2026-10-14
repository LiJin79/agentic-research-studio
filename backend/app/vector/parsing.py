import csv
import io
import json
from pathlib import Path

from docx import Document as WordDocument
from pptx import Presentation
from pypdf import PdfReader


class UnsupportedDocumentError(ValueError):
    """表示上传文件扩展名不在项目允许范围内。"""


def extract_text(filename: str, content: bytes) -> str:
    """提取已有文本与 Office 文档内容；扫描 PDF 不伪装成已完成 OCR。"""
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8-sig")
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    if suffix == ".json":
        value = json.loads(content.decode("utf-8-sig"))
        return json.dumps(value, ensure_ascii=False, indent=2)
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if suffix == ".docx":
        document = WordDocument(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join(paragraphs + tables)
    if suffix == ".pptx":
        presentation = Presentation(io.BytesIO(content))
        parts = []
        for number, slide in enumerate(presentation.slides, 1):
            parts.append(f"Slide {number}")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend(
                        " | ".join(cell.text for cell in row.cells) for row in shape.table.rows
                    )
        return "\n".join(parts)
    raise UnsupportedDocumentError("仅支持 txt、md、csv、json、pdf、docx 和 pptx 文件")
